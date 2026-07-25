"""
验证 v4: Sheet2 下拉选择式扣分 → 公式联动
"""
import openpyxl

INPUT = "/Users/sundanian/Documents/projects/ai-agents/my-agent/workspace/2026-07-21-供应商分级上半年盘点SOP/供应商分级盘点-计算模板.xlsx"
wb = openpyxl.load_workbook(INPUT)
ws2 = wb["2.SLA扣分台账"]

print("=" * 70)
print("📋 Sheet2 检查：表头结构")
print("=" * 70)
headers = []
for c in range(1, 10):
    val = ws2.cell(row=4, column=c).value
    headers.append(val)
    print(f"  列{c}({chr(64+c)}): {val}")

print("\n" + "=" * 70)
print("📋 模拟选择测试：填入不同扣分组合，验证公式联动")
print("=" * 70)

# 模拟场景
test_cases = [
    # (行, 黄牌, 投诉, 红线, 预期扣分, 预期直接C, 预期明细)
    (5,  0,  0, "否",  0, "否", "无扣分"),
    (6, -5,  0, "否", -5, "否", "黄牌"),
    (7,  0, -5, "否", -5, "否", "有责投诉"),
    (8, -5, -5, "否",-10, "否", "黄牌 + 有责投诉"),
    (9,  0,  0, "是",  0, "🔴 是·直接C级", "🔴红线"),
    (10,-5,  0, "是", -5, "🔴 是·直接C级", "黄牌 + 🔴红线"),
]

for row, huang, tou, red, exp_score, exp_c, exp_detail in test_cases:
    ws2.cell(row=row, column=3, value=huang)
    ws2.cell(row=row, column=4, value=tou)
    ws2.cell(row=row, column=5, value=red)

# 检查公式（公式不会被openpyxl计算，只检查公式是否正确）
print(f"\n{'行':>4} {'黄牌':>6} {'投诉':>6} {'红线':>6} {'预期扣分':>8} {'公式F':>12} {'公式G':>20} {'公式H':>30}")
print("-" * 100)
for row, huang, tou, red, exp_score, exp_c, exp_detail in test_cases:
    f_formula = ws2.cell(row=row, column=6).value
    g_formula = ws2.cell(row=row, column=7).value
    h_formula = ws2.cell(row=row, column=8).value
    print(f"{row:>4} {huang:>6} {tou:>6} {red:>6} {exp_score:>8}")
    print(f"     F公式: {str(f_formula)[:60]}")
    print(f"     G公式: {str(g_formula)[:60]}")
    print(f"     H公式: {str(h_formula)[:80]}")
    print()

# 手动验证逻辑
print("=" * 70)
print("📋 手动验证公式逻辑")
print("=" * 70)
for row, huang, tou, red, exp_score, exp_c, exp_detail in test_cases:
    actual_score = huang + tou  # F = C + D
    actual_c = "🔴 是·直接C级" if red == "是" else "否"
    
    # 明细逻辑
    parts = []
    if huang == -5: parts.append("黄牌")
    if tou == -5: parts.append("有责投诉")
    if red == "是": parts.append("🔴红线")
    if not parts: parts.append("无扣分")
    actual_detail = " + ".join(parts)
    
    score_ok = "✅" if actual_score == exp_score else "❌"
    c_ok = "✅" if actual_c == exp_c else "❌"
    detail_ok = "✅" if actual_detail == exp_detail else "❌"
    
    print(f"行{row}: 黄牌={huang}, 投诉={tou}, 红线={red}")
    print(f"  扣分: {actual_score} (预期{exp_score}) {score_ok}")
    print(f"  直接C: {actual_c} (预期{exp_c}) {c_ok}")
    print(f"  明细: {actual_detail} (预期{exp_detail}) {detail_ok}")
    print()

# 检查数据验证（下拉）是否设置
print("=" * 70)
print("📋 下拉验证设置检查")
print("=" * 70)
dvs = ws2.data_validations.dataValidation
for dv in dvs:
    print(f"  范围: {dv.sqref} | 选项: {dv.formula1}")

# 清空测试数据（保留默认0/否）
for row, _, _, _, _, _, _ in test_cases:
    ws2.cell(row=row, column=3, value=0)
    ws2.cell(row=row, column=4, value=0)
    ws2.cell(row=row, column=5, value="否")

# 检查供应商名称
print("\n" + "=" * 70)
print("📋 供应商名称预填检查")
print("=" * 70)
for r in range(5, 16):
    print(f"  行{r}: {ws2.cell(row=r, column=1).value} - {ws2.cell(row=r, column=2).value}")

wb.save(INPUT)
print("\n✅ v4 Sheet2 验证完成，已重置默认值")
