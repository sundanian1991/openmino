"""
验证Excel模板 v3：11家供应商 + 首贷/复贷 + 打分标准
"""
import openpyxl

INPUT = "/Users/sundanian/Documents/projects/ai-agents/my-agent/workspace/2026-07-21-供应商分级上半年盘点SOP/供应商分级盘点-计算模板.xlsx"

wb = openpyxl.load_workbook(INPUT)

# 检查供应商名称是否正确预填
ws1 = wb["1.赛马分原始数据"]
print("=" * 60)
print("📋 检查11家供应商名称预填")
print("=" * 60)
for r in range(6, 17):
    sid = ws1.cell(row=r, column=1).value
    name = ws1.cell(row=r, column=2).value
    print(f"  行{r}: {sid} - {name}")

# 检查Sheet3稳定性默认B档
ws3 = wb["3.稳定性检查表"]
print("\n" + "=" * 60)
print("📋 检查Sheet3稳定性默认B档")
print("=" * 60)
b_count = 0
for r in range(6, 17):
    for c in range(3, 15):
        val = ws3.cell(row=r, column=c).value
        if val == "B":
            b_count += 1
print(f"  预填B档格子数: {b_count}（应为 11家×3月×4项=132）")
print(f"  {'✅ 正确' if b_count == 132 else '❌ 错误'}")

# 检查打分标准是否写入
print("\n" + "=" * 60)
print("📋 检查稳定性检查表打分标准")
print("=" * 60)
for r in range(19, 25):
    dim = ws3.cell(row=r, column=1).value
    if dim:
        a_desc = ws3.cell(row=r, column=3).value
        print(f"  行{r}: {dim}")
        if a_desc:
            print(f"    A档: {a_desc[:40]}...")

# 检查Sheet4供管自评打分标准
ws4 = wb["4.供管自评表"]
print("\n" + "=" * 60)
print("📋 检查供管自评打分标准")
print("=" * 60)
for r in range(18, 24):
    dim = ws4.cell(row=r, column=1).value
    if dim:
        high_desc = ws4.cell(row=r, column=2).value
        print(f"  行{r}: {dim}")
        if high_desc:
            print(f"    高: {high_desc[:40]}...")

# 检查Sheet5 X档规则
ws5 = wb["5.评级结果"]
print("\n" + "=" * 60)
print("📋 检查Sheet5 X档规则（11家）")
print("=" * 60)
sample_formula = ws5.cell(row=5, column=9).value
print(f"  X档公式: {sample_formula}")
if "<=3" in str(sample_formula) and "<=8" in str(sample_formula):
    print("  ✅ 正确（前3高/4-8中/9-11低）")
else:
    print("  ❌ 需要检查")

print("\n" + "=" * 60)
print("✅ v3 模板验证完成")
print("=" * 60)
